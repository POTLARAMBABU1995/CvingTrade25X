from __future__ import annotations

import csv
import io
import os
import re
import threading
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import xlrd

from services.nse_symbol_normalization import canonical_nse_symbol, format_fyers_nse_symbol


DEFAULT_FYERS_NIFTY500_EXISTING_CSV = r"D:\fyers_api_integration\data\symbols_nifty500.csv"
FYERS_NIFTY500_EXISTING_CSV = (
    os.getenv("FYERS_NIFTY500_EXISTING_CSV", DEFAULT_FYERS_NIFTY500_EXISTING_CSV).strip()
    or DEFAULT_FYERS_NIFTY500_EXISTING_CSV
)
_HEADER_CANDIDATES = ("symbol", "symbols", "ticker", "tradingsymbol")
_HEADER_NORMALIZE_RE = re.compile(r"[^a-z0-9]+")
_SYMBOL_SANITIZE_RE = re.compile(r"[^A-Z0-9&._-]+")
_SYMBOL_SPLIT_RE = re.compile(r"[\s,;]+")
_NON_ALNUM_RE = re.compile(r"[^A-Z0-9]+")
_WRITE_LOCK = threading.Lock()
_SUPPORTED_UPLOAD_SUFFIXES = {".csv", ".xls", ".xlsx"}


def _resolve_existing_csv_path() -> Path:
    return Path(FYERS_NIFTY500_EXISTING_CSV).expanduser()


def resolve_existing_csv_path() -> Path:
    return _resolve_existing_csv_path()


def _normalize_symbol(value: object) -> str:
    cleaned = canonical_nse_symbol(str(value or "").strip().strip('"').strip("'"))
    canonical = _NON_ALNUM_RE.sub("", cleaned)
    if canonical in ("NIFTY500", "SYMBOL"):
        return ""
    return cleaned


def _format_existing_symbol(symbol: str) -> str:
    normalized = _normalize_symbol(symbol)
    if not normalized:
        raise ValueError("Symbol is required.")
    return format_fyers_nse_symbol(symbol)


def _dedupe_symbols(symbols: list[str]) -> list[str]:
    ordered: list[str] = []
    seen: set[str] = set()
    for symbol in symbols or []:
        normalized = _normalize_symbol(symbol)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        ordered.append(normalized)
    return ordered


def _extract_symbols_from_payload(payload: dict[str, Any]) -> list[str]:
    symbols_raw: list[Any] = []
    if isinstance(payload.get("symbols"), (list, tuple, set)):
        symbols_raw.extend(list(payload.get("symbols") or []))
    elif payload.get("symbols") is not None:
        symbols_raw.extend(_SYMBOL_SPLIT_RE.split(str(payload.get("symbols") or "")))

    if payload.get("symbol") is not None:
        symbols_raw.append(payload.get("symbol"))

    normalized: list[str] = []
    seen: set[str] = set()
    for raw in symbols_raw:
        token = _normalize_symbol(raw)
        if not token or token in seen:
            continue
        seen.add(token)
        normalized.append(token)
    return normalized


def _normalize_header_token(value: object) -> str:
    return _HEADER_NORMALIZE_RE.sub("", str(value or "").strip().lower())


def _resolve_symbol_index(header_row: list[str]) -> int | None:
    for idx, raw in enumerate(header_row):
        token = _normalize_header_token(raw)
        if (
            token in _HEADER_CANDIDATES
            or token.endswith("symbol")
            or token.endswith("symbols")
            or token.endswith("ticker")
        ):
            return idx
    return None


def _find_symbol_column(rows: list[list[Any]]) -> tuple[int, int] | None:
    for row_index, row in enumerate(rows):
        symbol_index = _resolve_symbol_index([str(value or "") for value in row])
        if symbol_index is not None:
            return row_index, symbol_index
    return None


def _read_symbols_from_text(csv_text: str) -> list[str]:
    text = str(csv_text or "").lstrip("\ufeff")
    if not text.strip():
        return []

    reader = csv.reader(io.StringIO(text), skipinitialspace=True)
    rows = [row for row in reader if row]
    if not rows:
        return []

    symbol_column = _find_symbol_column(rows)
    if symbol_column is None:
        symbol_index = 0
        start_index = 0
    else:
        start_index, symbol_index = symbol_column
        start_index += 1

    symbols: list[str] = []
    for row in rows[start_index:]:
        if symbol_index >= len(row):
            continue
        symbol = _normalize_symbol(row[symbol_index])
        if symbol:
            symbols.append(symbol)
    return _dedupe_symbols(symbols)


def _read_symbols_from_path(path: Path) -> list[str]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return _dedupe_symbols(_read_symbols_from_text(handle.read()))


def _rows_to_symbols(rows: list[list[Any]]) -> list[str]:
    if not rows:
        return []

    symbol_column = _find_symbol_column(rows)
    if symbol_column is None:
        return []
    header_row_index, symbol_index = symbol_column

    symbols: list[str] = []
    for row in rows[header_row_index + 1:]:
        if symbol_index >= len(row):
            continue
        symbol = _normalize_symbol(row[symbol_index])
        if symbol:
            symbols.append(symbol)
    return _dedupe_symbols(symbols)


def _decode_csv_upload(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Unable to decode the uploaded CSV file. Use UTF-8, ANSI, XLSX, or XLS.")


def _read_symbols_from_csv_bytes(content: bytes) -> list[str]:
    return _read_symbols_from_text(_decode_csv_upload(content))


def _read_symbols_from_xlsx_bytes(content: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    symbols: list[str] = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        symbols.extend(_rows_to_symbols(rows))
    workbook.close()
    return _dedupe_symbols(symbols)


def _read_symbols_from_xls_bytes(content: bytes) -> list[str]:
    workbook = xlrd.open_workbook(file_contents=content)
    symbols: list[str] = []
    for sheet in workbook.sheets():
        rows = [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
        symbols.extend(_rows_to_symbols(rows))
    return _dedupe_symbols(symbols)


def _read_symbols_from_uploaded_file(filename: str, content: bytes) -> list[str]:
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix not in _SUPPORTED_UPLOAD_SUFFIXES:
        raise ValueError("Only CSV, XLSX, and XLS files are supported for NIFTY500 sync uploads.")
    if suffix == ".csv":
        return _read_symbols_from_csv_bytes(content)
    if suffix == ".xlsx":
        return _read_symbols_from_xlsx_bytes(content)
    return _read_symbols_from_xls_bytes(content)


def _resolve_uploaded_symbols(payload: dict[str, Any]) -> tuple[list[str], str]:
    uploaded_files = payload.get("uploadedFiles")
    if not isinstance(uploaded_files, list) or not uploaded_files:
        return [], ""

    symbols: list[str] = []
    filenames: list[str] = []
    for item in uploaded_files:
        if not isinstance(item, dict):
            continue
        filename = str(item.get("filename") or "").strip()
        content = item.get("content")
        if not filename or not isinstance(content, (bytes, bytearray)):
            continue
        filenames.append(filename)
        symbols.extend(_read_symbols_from_uploaded_file(filename, bytes(content)))

    label = filenames[0] if len(filenames) == 1 else f"{len(filenames)} files uploaded"
    return _dedupe_symbols(symbols), label


def load_existing_symbols(*, require_file: bool = False) -> list[str]:
    path = _resolve_existing_csv_path()
    if require_file and not path.exists():
        raise FileNotFoundError(f"NIFTY500 universe file is missing: {path}")
    return _read_symbols_from_path(path)


def _summarize_symbols(symbols: list[str]) -> tuple[list[dict[str, Any]], dict[str, int]]:
    counts: dict[str, int] = {}
    ordered: list[str] = []
    for symbol in symbols:
        normalized = _normalize_symbol(symbol)
        if not normalized:
            continue
        if normalized not in counts:
            counts[normalized] = 0
            ordered.append(normalized)
        counts[normalized] += 1

    rows = [
        {
            "sno": idx + 1,
            "symbol": symbol,
            "symbolCount": counts[symbol],
        }
        for idx, symbol in enumerate(ordered)
    ]
    return rows, counts


def _decorate_rows(
    rows: list[dict[str, Any]],
    other_symbols: set[str] | None,
    role: str,
) -> list[dict[str, Any]]:
    decorated: list[dict[str, Any]] = []
    for row in rows:
        symbol = str(row.get("symbol") or "")
        matched = other_symbols is not None and symbol in other_symbols
        if other_symbols is None:
            flag = "-"
            match_state = "uncompared"
            recommended_action = "Read"
        elif matched:
            flag = "Y"
            match_state = "matched"
            recommended_action = "Read"
        elif role == "new":
            flag = "N"
            match_state = "missing_in_existing"
            recommended_action = "Create"
        else:
            flag = "N"
            match_state = "missing_in_new"
            recommended_action = "Read"

        decorated.append(
            {
                **row,
                "flag": flag,
                "matchState": match_state,
                "recommendedAction": recommended_action,
            }
        )
    return decorated


def _build_payload(
    *,
    existing_path: Path,
    existing_symbols: list[str],
    new_symbols: list[str] | None = None,
    filename: str | None = None,
    message: str | None = None,
    merged_count: int = 0,
) -> dict[str, Any]:
    existing_rows_base, existing_counts = _summarize_symbols(existing_symbols)
    new_rows_base, new_counts = _summarize_symbols(new_symbols or [])

    new_symbol_set = set(new_counts)
    existing_symbol_set = set(existing_counts)

    existing_rows = _decorate_rows(
        existing_rows_base,
        new_symbol_set if new_symbols is not None else None,
        "existing",
    )
    new_rows = _decorate_rows(
        new_rows_base,
        existing_symbol_set if new_symbols is not None else None,
        "new",
    )

    matched_symbols = sorted(existing_symbol_set & new_symbol_set)
    missing_in_existing = [
        row["symbol"] for row in new_rows if row.get("matchState") == "missing_in_existing"
    ]
    existing_only = [
        row["symbol"] for row in existing_rows if row.get("matchState") == "missing_in_new"
    ]

    return {
        "ok": True,
        "message": message or "NIFTY500 sync snapshot loaded.",
        "existing": {
            "path": str(existing_path),
            "exists": existing_path.exists(),
            "rows": existing_rows,
            "uniqueCount": len(existing_rows_base),
            "totalCount": sum(existing_counts.values()),
        },
        "new": {
            "filename": filename or "",
            "rows": new_rows,
            "uniqueCount": len(new_rows_base),
            "totalCount": sum(new_counts.values()),
            "loaded": new_symbols is not None,
        },
        "summary": {
            "matchedCount": len(matched_symbols),
            "missingInExistingCount": len(missing_in_existing),
            "existingOnlyCount": len(existing_only),
            "mergedCount": int(merged_count or 0),
            "requiresCreate": not existing_path.exists(),
        },
        "actions": {
            "matchedSymbols": matched_symbols,
            "missingInExisting": missing_in_existing,
            "existingOnly": existing_only,
        },
    }


def _require_uploaded_symbols(payload: dict[str, Any]) -> tuple[list[str], str]:
    uploaded_symbols, uploaded_label = _resolve_uploaded_symbols(payload)
    if uploaded_symbols:
        return uploaded_symbols, uploaded_label

    csv_text = str(payload.get("newCsvContent") or payload.get("csvText") or "")
    if not csv_text.strip():
        raise ValueError("Upload at least one CSV, XLSX, or XLS file that contains a symbol column.")
    filename = str(payload.get("filename") or payload.get("name") or "").strip()
    return _read_symbols_from_text(csv_text), filename


def _write_existing_symbols(path: Path, symbols: list[str]) -> None:
    deduped_symbols = _dedupe_symbols(symbols)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(f"{path.name}.{uuid.uuid4().hex}.tmp")
    with temp_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["symbol"])
        for symbol in deduped_symbols:
            writer.writerow([_format_existing_symbol(symbol)])
    temp_path.replace(path)


def get_nifty500_sync_snapshot() -> dict[str, Any]:
    path = _resolve_existing_csv_path()
    existing_symbols = _read_symbols_from_path(path)
    message = "Existing FYERS NIFTY500 file loaded." if path.exists() else "Existing FYERS NIFTY500 file is missing."
    return _build_payload(existing_path=path, existing_symbols=existing_symbols, message=message)


def compare_nifty500_sync(payload: dict[str, Any]) -> dict[str, Any]:
    new_symbols, filename = _require_uploaded_symbols(payload)
    path = _resolve_existing_csv_path()
    existing_symbols = _read_symbols_from_path(path)
    if not new_symbols:
        raise ValueError("No symbols were found in the uploaded file. Ensure the file contains a symbol column.")
    return _build_payload(
        existing_path=path,
        existing_symbols=existing_symbols,
        new_symbols=new_symbols,
        filename=filename,
        message="NIFTY500 compare completed.",
    )


def merge_missing_nifty500_symbols(payload: dict[str, Any]) -> dict[str, Any]:
    new_symbols, filename = _require_uploaded_symbols(payload)
    path = _resolve_existing_csv_path()
    if not new_symbols:
        raise ValueError("No symbols were found in the uploaded file. Ensure the file contains a symbol column.")

    with _WRITE_LOCK:
        existing_symbols = _read_symbols_from_path(path)
        existing_set = set(existing_symbols)
        new_rows_base, _ = _summarize_symbols(new_symbols)
        merged_symbols = [row["symbol"] for row in new_rows_base if row["symbol"] not in existing_set]
        merged_count = len(merged_symbols)
        if merged_symbols:
            updated_symbols = [*existing_symbols, *merged_symbols]
            _write_existing_symbols(path, updated_symbols)
            existing_symbols = updated_symbols

    message = (
        f"Merged {merged_count} missing symbol{'s' if merged_count != 1 else ''} into the existing FYERS CSV."
        if merged_count
        else "No new symbols were missing from the existing FYERS CSV."
    )
    return _build_payload(
        existing_path=path,
        existing_symbols=existing_symbols,
        new_symbols=new_symbols,
        filename=filename,
        message=message,
        merged_count=merged_count,
    )


def add_nifty500_symbol(payload: dict[str, Any]) -> dict[str, Any]:
    normalized_symbols = _extract_symbols_from_payload(payload)
    if not normalized_symbols:
        path = _resolve_existing_csv_path()
        existing_symbols = _read_symbols_from_path(path)
        return _build_payload(
            existing_path=path,
            existing_symbols=existing_symbols,
            message="No valid symbols to add. Blocked/invalid symbols were skipped.",
        )

    path = _resolve_existing_csv_path()
    with _WRITE_LOCK:
        existing_symbols = _read_symbols_from_path(path)
        existing_set = set(existing_symbols)
        added_symbols: list[str] = []
        for symbol in normalized_symbols:
            if symbol in existing_set:
                continue
            existing_set.add(symbol)
            existing_symbols.append(symbol)
            added_symbols.append(symbol)
        if added_symbols:
            _write_existing_symbols(path, existing_symbols)

    requested_count = len(normalized_symbols)
    added_count = len(added_symbols)
    skipped_count = max(requested_count - added_count, 0)
    if requested_count == 1:
        message = f"Symbol {normalized_symbols[0]} is available in the existing FYERS CSV."
    else:
        message = (
            f"Processed {requested_count} symbols. Added {added_count} new symbol"
            f"{'s' if added_count != 1 else ''}, skipped {skipped_count} existing."
        )

    return _build_payload(
        existing_path=path,
        existing_symbols=existing_symbols,
        message=message,
    )


def update_nifty500_symbol(payload: dict[str, Any]) -> dict[str, Any]:
    original_symbol = _normalize_symbol(payload.get("originalSymbol"))
    updated_symbol = _normalize_symbol(
        payload.get("symbol")
        or payload.get("updatedSymbol")
        or payload.get("updated_symbol")
    )
    if not original_symbol:
        raise ValueError("Original symbol is required.")
    if not updated_symbol:
        raise ValueError("Updated symbol is required.")

    path = _resolve_existing_csv_path()
    with _WRITE_LOCK:
        existing_symbols = _read_symbols_from_path(path)
        existing_set = set(existing_symbols)
        if original_symbol not in existing_set:
            raise ValueError(f"Symbol {original_symbol} was not found in the existing FYERS CSV.")
        if updated_symbol != original_symbol and updated_symbol in existing_set:
            raise ValueError(f"Symbol {updated_symbol} is already available in the existing FYERS CSV.")
        rewritten = _dedupe_symbols([updated_symbol if symbol == original_symbol else symbol for symbol in existing_symbols])
        _write_existing_symbols(path, rewritten)

    return _build_payload(
        existing_path=path,
        existing_symbols=rewritten,
        message=f"Updated {original_symbol} to {updated_symbol} in the existing FYERS CSV.",
    )


def delete_nifty500_symbol(payload: dict[str, Any]) -> dict[str, Any]:
    normalized = _normalize_symbol(payload.get("symbol"))
    if not normalized:
        raise ValueError("Symbol is required.")

    path = _resolve_existing_csv_path()
    with _WRITE_LOCK:
        existing_symbols = _read_symbols_from_path(path)
        if normalized not in set(existing_symbols):
            return {
                **_build_payload(
                    existing_path=path,
                    existing_symbols=existing_symbols,
                    message=f"Symbol {normalized} was already absent from the existing FYERS CSV.",
                ),
                "deleted": False,
                "symbol": normalized,
            }
        rewritten = [symbol for symbol in existing_symbols if symbol != normalized]
        _write_existing_symbols(path, rewritten)

    return {
        **_build_payload(
            existing_path=path,
            existing_symbols=rewritten,
            message=f"Deleted {normalized} from the existing FYERS CSV.",
        ),
        "deleted": True,
        "symbol": normalized,
    }
